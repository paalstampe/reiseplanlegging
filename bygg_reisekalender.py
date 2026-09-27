# Bygger Reisekalender.xlsx (mock-up). Bruk: python3 bygg_reisekalender.py [utfil.xlsx]
# Krever openpyxl. Formler beregnes når filen åpnes i Excel.
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import reisedata as RD
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

F="Arial"
def f(**k): return Font(name=F, size=k.pop('size',10), **k)
BLUE=f(color="0000FF"); BOLD=f(bold=True); TITLE=f(bold=True,size=14); NORM=f()
GREEN=f(color="008000")
HDR=PatternFill("solid",fgColor="DDE4EE"); INP=PatternFill("solid",fgColor="FFF9D6")
thin=Side(style="thin",color="BFBFBF")

wb=Workbook()

# ---------------- Les meg ----------------
lm=wb.active; lm.title="Les meg"
rows=[
("Reisekalender – beste tid for reisemål",TITLE),
("Mock-up. Masterliste = «Reise - utenfor Europa.md». Klimatall er omtrentlige og verifiseres ved populering.",f(italic=True)),
("",None),
("Flyt",BOLD),
("1. Destinasjoner – én rad per reisemål (fra .md-filen). ID brukes som nøkkel i alle andre ark.",NORM),
("2. Måned – 12 rader per destinasjon med delscorer, T/N og begrunnelse. Samlet score beregnes.",NORM),
("3. Vinduer – ukesavgrensede justeringer (hendelser, dyreliv, helligdager). Legges oppå månedsscoren.",NORM),
("4. Kalender – generert heatmap destinasjon × ISO-uke. Marker egne tilgjengelige uker med «x» i rad 7.",NORM),
("",None),
("Konvensjoner",BOLD),
("Blå tekst / gul bakgrunn = input. Svart = formel. Ikke skriv i formelceller.",NORM),
("Uke → måned: en uke tilhører måneden torsdagen faller i (ISO-logikk). Uke 53 (f.eks. 2026) vises ikke.",NORM),
("Vindu med Fra > Til går over nyttår (f.eks. uke 49–12).",NORM),
("",None),
("Scoreskala (1–5)",BOLD),
("Klima: 5 = ideelt · 3 = akseptabelt · 1 = frarådes (monsun, syklon/tyfon, ekstrem hete/kulde)",NORM),
("Opplevelse: 5 = toppsesong for det man reiser for · 1 = stengt / ikke gjennomførbart",NORM),
("Trengsel/pris: 5 = lite folk, lave priser · 1 = maks høysesong, utsolgt",NORM),
("",None),
("Vekter for samlet score",BOLD),
]
for i,(t,ft) in enumerate(rows,1):
    c=lm.cell(i,2,t)
    if ft: c.font=ft
wr=len(rows)+1
weights=[("Klima",0.4),("Opplevelse",0.4),("Trengsel/pris",0.2)]
for j,(n,w) in enumerate(weights):
    lm.cell(wr+j,2,n).font=NORM
    c=lm.cell(wr+j,3,w); c.font=BLUE; c.fill=INP; c.number_format="0%"
lm.cell(wr+3,2,"Sum (skal være 100 %)").font=NORM
c=lm.cell(wr+3,3,f"=SUM(C{wr}:C{wr+2})"); c.number_format="0%"; c.font=NORM
lm.cell(wr+4,2,"Veto-tak: Hvis Klima eller Opplevelse = 1, begrenses samlet score til dette taket (vektet snitt hvis lavere).").font=NORM
c=lm.cell(wr+4,3,1.5); c.font=BLUE; c.fill=INP; c.number_format="0.0"
lm.cell(wr+5,2,"Vindu-justering (±) legges til i Kalender; resultatet begrenses til 1–5.").font=NORM
WK,WO,WT=f"'Les meg'!$C${wr}",f"'Les meg'!$C${wr+1}",f"'Les meg'!$C${wr+2}"
VT=f"'Les meg'!$C${wr+4}"
lm.column_dimensions['A'].width=2; lm.column_dimensions['B'].width=95; lm.column_dimensions['C'].width=10
lm.sheet_view.showGridLines=False

# ---------------- Destinasjoner ----------------
d=wb.create_sheet("Destinasjoner")
d["A1"]="Destinasjoner (master: .md-filen)"; d["A1"].font=TITLE
hdr=["ID","Region","Land","Sted / fokus","Visningsnavn","Status","Kategori","Varighet (dager)","Data i Måned","Merknad","Klimakilde (stasjon, periode)"]
for j,h in enumerate(hdr,1):
    c=d.cell(4,j,h); c.font=BOLD; c.fill=HDR
AS="Asia, Oseania og Midt-Østen"; NA="Nord-Amerika"; SA="Sør- og Mellom-Amerika"; AF="Afrika"
dest=[
("BTN",AS,"Bhutan","","","","",""),
("MMR",AS,"Myanmar","","","","",""),
("LAO",AS,"Laos","","","","",""),
("NPL",AS,"Nepal","Katmandu, Pokhara","","","",""),
("IND-GOA",AS,"India","Goa","","","",""),
("LKA-SV",AS,"Sri Lanka","Vest- og sørkysten, høylandet","","","","Sørvestmonsun mai–sep; best des–apr"),
("LKA-O",AS,"Sri Lanka","Østkysten (Trincomalee, Arugam Bay)","","","","Nordøstmonsun okt–jan; best mai–sep"),
("CHN",AS,"Kina","Beijing, Shanghai","","","",""),
("PYF",AS,"Fransk Polynesia","","","","",""),
("YEM-SOC",AS,"Yemen","Socotra","","","",""),
("VNM",AS,"Vietnam","Reunification Express + Hue","Ikke besøkt","Kultur / by / tog",14,"Hele landet langs ruten; Hue (sentral-Vietnam) er begrensende faktor for klima"),
("SILK",AS,"Sentral-Asia","Silkeveien","","","","Rute avklares (Usbekistan, Kirgisistan, Kasakhstan, Tadsjikistan, Turkmenistan)"),
("IDN-ROT",AS,"Indonesia","Rote","","","",""),
("SGP",AS,"Singapore","","","","",""),
("PHL",AS,"Filippinene","","","","",""),
("PNG",AS,"Papua Ny-Guinea","","","","",""),
("TLS",AS,"Øst-Timor","","","","",""),
("LBN",AS,"Libanon","","","","",""),
("USA-WC",NA,"USA Nordvestkyst","San Francisco til Vancouver","","","",""),
("USA-HI",NA,"Hawaii","","","","",""),
("USA-CHS",NA,"USA","Savannah-Charleston","","","",""),
("CRI",SA,"Costa Rica","","","","",""),
("PAN",SA,"Panama","","","","",""),
("COL",SA,"Colombia","","","","",""),
("DOM",SA,"Den dominikanske republikk","","","","",""),
("BRA-NE",SA,"Brasil","Rio og nordover","","","",""),
("ECU-GAL",SA,"Ecuador","Galápagos","","","",""),
("ATA",SA,"Antarktis","Antarktishalvøya (via Ushuaia)","Ikke besøkt","Natur / ekspedisjon",11,"Kun ekspedisjonscruise nov–mars"),
("MEX-CDMX",SA,"Mexico City","","","","",""),
("RWA",AF,"Rwanda","","","","",""),
("GAB",AF,"Gabon","","","","",""),
("GHA",AF,"Ghana","","","","",""),
("ETH",AF,"Etiopia","","","","",""),
("BWA",AF,"Botswana","Okavango, Chobe","Ikke besøkt","Safari",10,"Evt. kombinert med Zambia/Zimbabwe (Victoria Falls)"),
("ZMB",AF,"Zambia","","","","",""),
("ZWE",AF,"Zimbabwe","","","","",""),
("STP",AF,"São Tomé og Príncipe","","","","",""),
("EGY",AF,"Egypt","Kairo, Luxor","","","",""),
("BEN",AF,"Benin","","","","",""),
("TGO",AF,"Togo","","","","",""),
]
r0=5
for i,row in enumerate(dest):
    r=r0+i
    idd,reg,land,sted,stat,kat,var,mer=row
    kat=kat or RD.KATEGORI.get(idd,""); mer=mer or RD.MERKNAD.get(idd,"")
    kilde=RD.MANED[idd][0] if idd in RD.MANED else None
    vals={1:idd,2:reg,3:land,4:sted or None,6:stat or "Ikke besøkt",7:kat or None,8:var or None,10:mer or None,11:kilde}
    for col,v in vals.items():
        c=d.cell(r,col,v); c.font=BLUE
    d.cell(r,5,f'=IF(A{r}="","",C{r}&IF(D{r}<>""," – "&D{r},""))').font=NORM
    d.cell(r,9,f'=IF(A{r}="","",IF(COUNTIF(Måned!$A$5:$A$2000,A{r})>0,"Ja",""))').font=NORM
last_dest=r0+len(dest)-1
for r in range(last_dest+1,last_dest+21):  # spare rows
    d.cell(r,5,f'=IF(A{r}="","",C{r}&IF(D{r}<>""," – "&D{r},""))').font=NORM
    d.cell(r,9,f'=IF(A{r}="","",IF(COUNTIF(Måned!$A$5:$A$2000,A{r})>0,"Ja",""))').font=NORM
DEND=last_dest+20
dv=DataValidation(type="list",formula1='"Ikke besøkt,Planlagt,Booket,Besøkt"',allow_blank=True)
d.add_data_validation(dv); dv.add(f"F5:F{DEND}")
for col,w in zip("ABCDEFGHIJK",[10,26,24,30,40,12,18,10,10,50,60]): d.column_dimensions[col].width=w
d.freeze_panes="B5"; d.auto_filter.ref=f"A4:K{DEND}"
d["A2"]="Legg nye destinasjoner nederst (reserverte rader har formler). ID må være unik.";d["A2"].font=f(italic=True)

# ---------------- Måned ----------------
m=wb.create_sheet("Måned")
m["A1"]="Månedsvurdering"; m["A1"].font=TITLE
m["A2"]="T maks/nedbør = månedsnormaler fra én representativ stasjon (se Destinasjoner, kolonne K). Botswana og Antarktis er foreløpig ikke verifisert."; m["A2"].font=f(italic=True)
mh=["ID","Destinasjon","Mnd","Måned","Klima","Opplevelse","Trengsel/pris","Samlet","T maks °C","Nedbør mm","Begrunnelse"]
for j,h in enumerate(mh,1):
    c=m.cell(4,j,h); c.font=BOLD; c.fill=HDR
mnames=["jan","feb","mar","apr","mai","jun","jul","aug","sep","okt","nov","des"]
r=5
for idd in [x[0] for x in dest if x[0] in RD.MANED]:
    for mi,row in enumerate(RD.MANED[idd][1],1):
        k,o,t,T,N,txt=row
        m.cell(r,1,idd).font=BLUE
        m.cell(r,2,f'=IFERROR(INDEX(Destinasjoner!$E$5:$E${DEND},MATCH(A{r},Destinasjoner!$A$5:$A${DEND},0)),"")').font=NORM
        m.cell(r,3,mi).font=BLUE
        m.cell(r,4,f'=IF(C{r}="","",CHOOSE(C{r},"jan","feb","mar","apr","mai","jun","jul","aug","sep","okt","nov","des"))').font=NORM
        for col,v in zip([5,6,7,9,10,11],[k,o,t,T,N,txt]):
            c=m.cell(r,col,v); c.font=BLUE
        m.cell(r,8,f'=IF(COUNT(E{r}:G{r})<3,"",ROUND(IF(OR(E{r}=1,F{r}=1),MIN({VT},E{r}*{WK}+F{r}*{WO}+G{r}*{WT}),E{r}*{WK}+F{r}*{WO}+G{r}*{WT}),1))').font=NORM
        m.cell(r,8).number_format="0.0"
        r+=1
MEND=r-1
dvs=DataValidation(type="whole",operator="between",formula1="1",formula2="5",allow_blank=True)
m.add_data_validation(dvs); dvs.add("E5:G2000")
m.conditional_formatting.add("E5:H2000",ColorScaleRule(start_type="num",start_value=1,start_color="F8696B",mid_type="num",mid_value=3,mid_color="FFEB84",end_type="num",end_value=5,end_color="63BE7B"))
for col,w in zip("ABCDEFGHIJK",[10,40,6,7,8,11,13,9,10,11,70]): m.column_dimensions[col].width=w
m.freeze_panes="C5"; m.auto_filter.ref=f"A4:K{MEND}"

# ---------------- Vinduer ----------------
v=wb.create_sheet("Vinduer")
v["A1"]="Vinduer (ukesavgrensede justeringer)"; v["A1"].font=TITLE
v["A2"]="Fra > Til = går over nyttår. «Kun år» tom = gjelder hvert år (for bevegelige høytider: én rad per år)."; v["A2"].font=f(italic=True)
vh=["ID","Destinasjon","Fra uke","Til uke","Type","Justering (±)","Kun år","Beskrivelse"]
for j,h in enumerate(vh,1):
    c=v.cell(4,j,h); c.font=BOLD; c.fill=HDR
win=RD.VINDUER
for i,row in enumerate(win):
    r=5+i
    for col,val in zip([1,3,4,5,6,7,8],row):
        c=v.cell(r,col,val); c.font=BLUE
    v.cell(r,2,f'=IFERROR(INDEX(Destinasjoner!$E$5:$E${DEND},MATCH(A{r},Destinasjoner!$A$5:$A${DEND},0)),"")').font=NORM
for r in range(5+len(win),60):
    v.cell(r,2,f'=IF(A{r}="","",IFERROR(INDEX(Destinasjoner!$E$5:$E${DEND},MATCH(A{r},Destinasjoner!$A$5:$A${DEND},0)),"?"))').font=NORM
dvt=DataValidation(type="list",formula1='"Klima,Dyreliv,Hendelse,Helligdag,Stengt,Annet"',allow_blank=True)
v.add_data_validation(dvt); dvt.add("E5:E300")
for col,w in zip("ABCDEFGH",[10,40,8,8,12,12,8,80]): v.column_dimensions[col].width=w
v.freeze_panes="C5"

# ---------------- Kalender ----------------
k=wb.create_sheet("Kalender")
k["A1"]="Kalender – score per ISO-uke"; k["A1"].font=TITLE
k["A2"]="År"; k["A2"].font=BOLD
k["B2"]=2026; k["B2"].font=BLUE; k["B2"].fill=INP
k["C2"]="Månedsscore + vinduer, begrenset til 1–5. Tomt = ikke vurdert ennå."; k["C2"].font=f(italic=True)
k["A3"]="Grønn fra"; k["A3"].font=BOLD
k["B3"]=4; k["B3"].font=BLUE; k["B3"].fill=INP; k["B3"].number_format="0.0"
k["C3"]="Uker med score ≥ denne verdien farges lysegrønne; resten er uten farge."; k["C3"].font=f(italic=True)
k["A4"]="Mørkegrønn fra"; k["A4"].font=BOLD
k["B4"]=4.4; k["B4"].font=BLUE; k["B4"].fill=INP; k["B4"].number_format="0.0"
k["C4"]="De beste ukene (score ≥ denne verdien) farges mørkegrønne."; k["C4"].font=f(italic=True)
labels={4:"Uke",5:"Mandag",6:"Måned",7:"Tilgjengelig (x)"}
for rr,t in labels.items():
    c=k.cell(rr,5,t); c.font=BOLD; c.alignment=Alignment(horizontal="right")
FC=6; NW=52; LC=FC+NW-1
for w in range(1,NW+1):
    col=FC+w-1; cl=L(col)
    k.cell(4,col,w).font=BOLD
    c=k.cell(5,col,f"=DATE($B$2,1,4)-WEEKDAY(DATE($B$2,1,4),2)+1+({cl}4-1)*7"); c.number_format="d.m"; c.font=f(size=8)
    mn=f'CHOOSE(MONTH({cl}5+3),"jan","feb","mar","apr","mai","jun","jul","aug","sep","okt","nov","des")'
    if w==1: k.cell(6,col,"="+mn)
    else: k.cell(6,col,f'=IF(MONTH({cl}5+3)<>MONTH({L(col-1)}5+3),{mn},"")')
    k.cell(6,col).font=f(size=9,bold=True)
    c=k.cell(7,col, "x" if 45<=w<=48 else None); c.font=BLUE; c.fill=INP
    for rr in (4,5,6,7): k.cell(rr,col).alignment=Alignment(horizontal="center")
k.cell(7,FC).comment=Comment("Eksempel: uke 45–48 (november) er markert. Endre fritt.","Claude")
kh=["ID","Destinasjon","Region","Snitt mine uker","Laveste mine uker"]
for j,h in enumerate(kh,1):
    c=k.cell(8,j,h); c.font=BOLD; c.fill=HDR; c.alignment=Alignment(wrap_text=True,vertical="center")
for col in range(FC,LC+1): k.cell(8,col).fill=HDR
MA=f"Måned!$A$5:$A$2000"; MC=f"Måned!$C$5:$C$2000"; MH=f"Måned!$H$5:$H$2000"
VA,VC,VD,VF,VG=[f"Vinduer!${c}$5:${c}$300" for c in "ACDFG"]
KR0=9
nrows=DEND-5+1
for i in range(nrows):
    r=KR0+i; dr=5+i
    k.cell(r,1,f'=IF(Destinasjoner!A{dr}="","",Destinasjoner!A{dr})').font=NORM
    k.cell(r,2,f'=IF(A{r}="","",Destinasjoner!E{dr})').font=NORM
    k.cell(r,3,f'=IF(A{r}="","",Destinasjoner!B{dr})').font=NORM
    rng=f"{L(FC)}{r}:{L(LC)}{r}"; av=f"${L(FC)}$7:${L(LC)}$7"
    c=k.cell(r,4,f'=IFERROR(SUMIFS({rng},{av},"x")/COUNTIFS({av},"x",{rng},">0"),"")'); c.number_format="0.0"; c.font=BOLD
    c=k.cell(r,5,f'=IF(D{r}="","",_xlfn.MINIFS({rng},{av},"x",{rng},">0"))'); c.number_format="0.0"; c.font=NORM
    for w in range(1,NW+1):
        col=FC+w-1; cl=L(col); mon=f"MONTH({cl}$5+3)"
        inwin=(f"(({VC}<={VD})*({cl}$4>={VC})*({cl}$4<={VD})"
               f"+({VC}>{VD})*((({cl}$4>={VC})+({cl}$4<={VD}))>0))")
        adj=f"SUMPRODUCT(({VA}=$A{r})*{VF}*((({VG}=\"\")+({VG}=$B$2))>0)*{inwin})"
        fml=(f'=IF($A{r}="","",IF(COUNTIFS({MA},$A{r},{MC},{mon})=0,"",'
             f'MAX(1,MIN(5,SUMIFS({MH},{MA},$A{r},{MC},{mon})+{adj}))))')
        c=k.cell(r,col,fml); c.number_format="0.0"; c.font=f(size=8); c.alignment=Alignment(horizontal="center")
KEND=KR0+nrows-1
c=f"{L(FC)}{KR0}"; k.conditional_formatting.add(f"{L(FC)}{KR0}:{L(LC)}{KEND}",FormulaRule(formula=[f'AND(ISNUMBER({c}),{c}>=$B$4-0.001)'],fill=PatternFill(fill_type="solid",start_color="FF1E7B45",end_color="FF1E7B45"),font=Font(bold=True,color="FFFFFFFF"),stopIfTrue=True)); k.conditional_formatting.add(f"{L(FC)}{KR0}:{L(LC)}{KEND}",FormulaRule(formula=[f'AND(ISNUMBER({c}),{c}>=$B$3-0.001)'],fill=PatternFill(fill_type="solid",start_color="FF63BE7B",end_color="FF63BE7B"),font=Font(bold=True)))
c=f"D{KR0}"; k.conditional_formatting.add(f"D{KR0}:E{KEND}",FormulaRule(formula=[f'AND(ISNUMBER({c}),{c}>=$B$4-0.001)'],fill=PatternFill(fill_type="solid",start_color="FF1E7B45",end_color="FF1E7B45"),font=Font(bold=True,color="FFFFFFFF"),stopIfTrue=True)); k.conditional_formatting.add(f"D{KR0}:E{KEND}",FormulaRule(formula=[f'AND(ISNUMBER({c}),{c}>=$B$3-0.001)'],fill=PatternFill(fill_type="solid",start_color="FF63BE7B",end_color="FF63BE7B"),font=Font(bold=True)))
k.conditional_formatting.add(f"{L(FC)}4:{L(LC)}6",FormulaRule(formula=[f'{L(FC)}$7="x"'],fill=PatternFill(fill_type="solid",start_color="FF9BC2E6",end_color="FF9BC2E6"),font=Font(bold=True)))
# Månedsskille: kantlinje til venstre for første uke i hver måned (følger valgt år)
k.conditional_formatting.add(f"{L(FC)}4:{L(LC)}{KEND}",FormulaRule(formula=[f'IFERROR(MONTH({L(FC)}$5+3)<>MONTH({L(FC-1)}$5+3),FALSE)'],border=Border(left=Side(style="thin",color="FF000000"))))
k.column_dimensions['A'].width=14; k.column_dimensions['B'].width=38; k.column_dimensions['C'].width=22
k.column_dimensions['D'].width=9; k.column_dimensions['E'].width=13
for col in range(FC,LC+1): k.column_dimensions[L(col)].width=4.3
k.row_dimensions[8].height=28
k.freeze_panes=k.cell(KR0,FC)
k.auto_filter.ref=f"A8:{L(LC)}{KEND}"

wb.move_sheet("Kalender",offset=-3)
wb.active=1
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if c.font and c.font.name!=F: c.font=Font(name=F,size=c.font.size or 10,bold=c.font.bold,color=c.font.color,italic=c.font.italic)
out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "Reisekalender.xlsx")
wb.save(out)
print("Lagret:", out)
